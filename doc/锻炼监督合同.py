from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ============ 页面设置 ============
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3)
section.right_margin = Cm(3)

# ============ 样式辅助函数 ============
def set_paragraph_spacing(paragraph, before=0, after=0, line_spacing=1.5):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing

def add_run(paragraph, text, font_name='宋体', font_size=12, bold=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run

# ============ 合同标题 ============
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=0, after=6, line_spacing=1.5)
add_run(p, '锻炼监督协议书', '黑体', 22, bold=True)

# 副标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=0, after=12, line_spacing=1.0)
add_run(p, '（简易版）', '楷体', 14, color=(128, 128, 128))

# ============ 当事人信息 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=6, after=6, line_spacing=1.8)
add_run(p, '甲方（监督人）：', '宋体', 12, bold=True)
add_run(p, '彭红丽', '宋体', 12)

p = doc.add_paragraph()
set_paragraph_spacing(p, before=0, after=6, line_spacing=1.8)
add_run(p, '乙方（被监督人）：', '宋体', 12, bold=True)
add_run(p, '曾鸿发', '宋体', 12)

# 分隔线
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=6, after=6, line_spacing=1.0)
add_run(p, '━' * 30, '宋体', 8, color=(180, 180, 180))

# ============ 鉴于条款 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=6, after=6, line_spacing=1.8)
add_run(p, '鉴于', '宋体', 12, bold=True)
add_run(p, '甲方愿意对乙方的跑步锻炼进行监督，乙方自愿接受甲方的监督，双方本着诚实守信、相互督促的原则，经友好协商，就锻炼监督事宜达成如下协议：', '宋体', 12)

# ============ 第一条 锻炼安排 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=1.8)
add_run(p, '第一条  锻炼安排', '黑体', 12, bold=True)

items_1 = [
    '乙方应于每周一、周三、周五、周日进行跑步锻炼。',
    '每次锻炼时间：晚21:00开始，约22:00结束。',
    '乙方在每次开始跑步前（约21:00），须向甲方发送通知，告知"开始跑步"。',
    '乙方在每次跑步结束后（约22:00），须向甲方进行打卡汇报，确认当日锻炼已完成。',
]
for i, item in enumerate(items_1, 1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=2, line_spacing=1.8)
    p.paragraph_format.left_indent = Cm(0.75)
    add_run(p, f'（{i}）', '宋体', 12)
    add_run(p, item, '宋体', 12)

# ============ 第二条 监督义务 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=1.8)
add_run(p, '第二条  监督义务', '黑体', 12, bold=True)

items_2 = [
    '甲方有义务在约定锻炼时间前后提醒乙方进行跑步锻炼，但该提醒义务为非强制性，甲方未提醒不构成违约。',
    '甲方有权对乙方的锻炼执行情况进行监督，并在乙方未按时打卡时予以督促。',
    '甲方应以合理、善意的方式履行监督职责，尊重乙方的身体状况与个人意愿。',
]
for i, item in enumerate(items_2, 1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=2, line_spacing=1.8)
    p.paragraph_format.left_indent = Cm(0.75)
    add_run(p, f'（{i}）', '宋体', 12)
    add_run(p, item, '宋体', 12)

# ============ 第三条 罚则 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=1.8)
add_run(p, '第三条  未打卡罚则', '黑体', 12, bold=True)

items_3 = [
    '乙方在约定锻炼日跑步结束后，未向甲方打卡汇报的，每次须向甲方支付罚金人民币8元（大写：捌元整）。',
    '罚金由乙方主动向甲方支付，甲方有权进行追缴提醒。',
    '因不可抗力（如突发疾病、极端天气等）导致无法完成锻炼或打卡的，经及时告知甲方后，可免除当日罚金。',
]
for i, item in enumerate(items_3, 1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=2, line_spacing=1.8)
    p.paragraph_format.left_indent = Cm(0.75)
    add_run(p, f'（{i}）', '宋体', 12)
    add_run(p, item, '宋体', 12)

# ============ 第四条 监督报酬 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=1.8)
add_run(p, '第四条  监督报酬', '黑体', 12, bold=True)

items_4 = [
    '甲方履行监督职责的报酬为人民币20元/月（大写：贰拾元整），由乙方按月向甲方支付。',
    '报酬支付时间：每月最后一日之前。',
]
for i, item in enumerate(items_4, 1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=2, line_spacing=1.8)
    p.paragraph_format.left_indent = Cm(0.75)
    add_run(p, f'（{i}）', '宋体', 12)
    add_run(p, item, '宋体', 12)

# ============ 第五条 协议期限 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=1.8)
add_run(p, '第五条  协议期限与解除', '黑体', 12, bold=True)

items_5 = [
    '本协议自双方签字之日起生效，有效期至双方协商一致解除之日止。',
    '任何一方如需解除本协议，应提前7日书面通知对方。',
]
for i, item in enumerate(items_5, 1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=2, line_spacing=1.8)
    p.paragraph_format.left_indent = Cm(0.75)
    add_run(p, f'（{i}）', '宋体', 12)
    add_run(p, item, '宋体', 12)

# ============ 第六条 其他 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=1.8)
add_run(p, '第六条  其他约定', '黑体', 12, bold=True)

items_6 = [
    '本协议未尽事宜，由双方协商补充，补充内容与本协议具有同等效力。',
    '本协议一式两份，甲乙双方各执一份，具有同等法律效力。',
]
for i, item in enumerate(items_6, 1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=2, line_spacing=1.8)
    p.paragraph_format.left_indent = Cm(0.75)
    add_run(p, f'（{i}）', '宋体', 12)
    add_run(p, item, '宋体', 12)

# 分隔线
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph_spacing(p, before=12, after=12, line_spacing=1.0)
add_run(p, '━' * 30, '宋体', 8, color=(180, 180, 180))

# ============ 签字栏 ============
p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=2.0)
add_run(p, '甲方签字：', '宋体', 12, bold=True)
add_run(p, '________________', '宋体', 12)

p = doc.add_paragraph()
set_paragraph_spacing(p, before=6, after=18, line_spacing=2.0)
add_run(p, '乙方签字：', '宋体', 12, bold=True)
add_run(p, '________________', '宋体', 12)

p = doc.add_paragraph()
set_paragraph_spacing(p, before=12, after=6, line_spacing=2.0)
add_run(p, '签署日期：', '宋体', 12, bold=True)
add_run(p, '______年______月______日', '宋体', 12)

# ============ 保存 ============
output_path = os.path.join(os.path.expanduser('~'), 'Desktop', '锻炼监督协议书.docx')
doc.save(output_path)
print(f'文档已保存至：{output_path}')
